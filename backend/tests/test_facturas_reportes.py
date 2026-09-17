"""Pruebas de facturacion y de los reportes en PDF y Excel (quinto avance)."""

from datetime import date

from tests.conftest import cabecera

MOTO = {'tipo_item': 'producto', 'id_item': 1, 'cantidad': 1}

# Firmas de los formatos: un PDF empieza por %PDF y un .xlsx es un ZIP (PK).
FIRMA_PDF = b'%PDF'
FIRMA_XLSX = b'PK'
TIPO_EXCEL = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def vender(cliente_http, token, **extra):
    cuerpo = {'items': [MOTO], 'metodo_pago': 'efectivo', **extra}
    return cliente_http.post('/api/ventas', json=cuerpo, headers=cabecera(token)).json()


# ---------------------------------------------------------------------------
# Facturas
# ---------------------------------------------------------------------------
def test_la_factura_copia_los_datos_del_cliente(cliente_http, token_cliente):
    id_factura = vender(cliente_http, token_cliente)['id_factura']

    factura = cliente_http.get(
        f'/api/facturas/{id_factura}', headers=cabecera(token_cliente),
    ).json()['factura']

    assert factura['cliente_nombre'] == 'Carlos Perez'
    assert factura['cliente_documento'] == '1000000002'
    assert float(factura['porcentaje_iva']) == 19
    assert len(factura['detalles']) == 1


def test_una_venta_no_se_puede_facturar_dos_veces(cliente_http, token_cliente, token_admin):
    id_venta = vender(cliente_http, token_cliente)['id_venta']

    respuesta = cliente_http.post(
        '/api/facturas', json={'venta_id': id_venta}, headers=cabecera(token_admin),
    )

    assert respuesta.status_code == 409
    assert 'ya tiene la factura' in respuesta.json()['message']


def test_se_puede_vender_sin_emitir_factura_y_emitirla_despues(cliente_http, token_cliente, token_admin):
    venta = vender(cliente_http, token_cliente, generar_factura=False)
    assert venta['numero_factura'] is None

    respuesta = cliente_http.post(
        '/api/facturas', json={'venta_id': venta['id_venta']}, headers=cabecera(token_admin),
    )

    assert respuesta.status_code == 201, respuesta.text
    assert respuesta.json()['numero_factura'].startswith('FV-')


def test_la_factura_se_descarga_en_pdf(cliente_http, token_cliente):
    id_factura = vender(cliente_http, token_cliente)['id_factura']

    respuesta = cliente_http.get(
        f'/api/facturas/{id_factura}/pdf', headers=cabecera(token_cliente),
    )

    assert respuesta.status_code == 200
    assert respuesta.headers['content-type'] == 'application/pdf'
    assert respuesta.content.startswith(FIRMA_PDF)
    assert 'FV-' in respuesta.headers['content-disposition']


def test_el_cliente_no_puede_descargar_la_factura_de_otro(cliente_http, token_cliente, token_admin):
    ajena = vender(cliente_http, token_admin)['id_factura']

    respuesta = cliente_http.get(f'/api/facturas/{ajena}/pdf', headers=cabecera(token_cliente))

    assert respuesta.status_code == 403


def test_anular_la_factura_anula_la_venta_y_devuelve_el_stock(cliente_http, token_cliente, token_admin):
    antes = cliente_http.get('/api/productos/1').json()['producto']['stock']
    venta = vender(cliente_http, token_cliente)

    cliente_http.patch(
        f"/api/facturas/{venta['id_factura']}/estado",
        json={'estado': 'anulada'}, headers=cabecera(token_admin),
    )

    detalle = cliente_http.get(
        f"/api/ventas/{venta['id_venta']}", headers=cabecera(token_admin),
    ).json()['venta']
    assert detalle['estado'] == 'anulada'
    assert cliente_http.get('/api/productos/1').json()['producto']['stock'] == antes


def test_la_busqueda_de_facturas_encuentra_por_numero(cliente_http, token_cliente, token_admin):
    numero = vender(cliente_http, token_cliente)['numero_factura']

    encontradas = cliente_http.get(
        '/api/facturas', params={'busqueda': numero}, headers=cabecera(token_admin),
    ).json()['facturas']

    assert len(encontradas) == 1
    assert encontradas[0]['numero_factura'] == numero


# ---------------------------------------------------------------------------
# Reporte diario de ventas
# ---------------------------------------------------------------------------
def test_el_reporte_del_dia_trae_las_ventas_y_sus_totales(cliente_http, token_cliente, token_admin):
    vender(cliente_http, token_cliente)
    vender(cliente_http, token_cliente)

    reporte = cliente_http.get('/api/reportes/ventas-diarias', headers=cabecera(token_admin)).json()

    assert reporte['fecha'] == date.today().isoformat()
    assert reporte['resumen']['cantidad'] == 2
    assert reporte['resumen']['unidades'] == 2
    assert float(reporte['resumen']['total']) == 480000 * 1.19 * 2
    assert reporte['negocio']['nombre']


def test_el_reporte_de_un_dia_sin_ventas_sale_vacio_pero_valido(cliente_http, token_admin):
    reporte = cliente_http.get(
        '/api/reportes/ventas-diarias',
        params={'fecha': '2020-01-01'}, headers=cabecera(token_admin),
    ).json()

    assert reporte['ventas'] == []
    assert float(reporte['resumen']['total']) == 0


def test_el_reporte_se_exporta_a_pdf(cliente_http, token_cliente, token_admin):
    vender(cliente_http, token_cliente)

    respuesta = cliente_http.get('/api/reportes/ventas-diarias/pdf', headers=cabecera(token_admin))

    assert respuesta.status_code == 200
    assert respuesta.headers['content-type'] == 'application/pdf'
    assert respuesta.content.startswith(FIRMA_PDF)
    # Un PDF con contenido real pesa bastante mas que una cabecera suelta.
    assert len(respuesta.content) > 2000


def test_el_reporte_se_exporta_a_excel(cliente_http, token_cliente, token_admin):
    vender(cliente_http, token_cliente)

    respuesta = cliente_http.get('/api/reportes/ventas-diarias/excel', headers=cabecera(token_admin))

    assert respuesta.status_code == 200
    assert respuesta.headers['content-type'] == TIPO_EXCEL
    assert respuesta.content.startswith(FIRMA_XLSX)
    assert '.xlsx' in respuesta.headers['content-disposition']


def test_el_excel_trae_las_tres_hojas_con_los_datos(cliente_http, token_cliente, token_admin):
    from io import BytesIO

    from openpyxl import load_workbook

    vender(cliente_http, token_cliente)
    respuesta = cliente_http.get('/api/reportes/ventas-diarias/excel', headers=cabecera(token_admin))

    libro = load_workbook(BytesIO(respuesta.content))
    assert libro.sheetnames == ['Ventas', 'Detalle', 'Resumen']
    # Fila 4 es la cabecera y la 5 la primera venta.
    assert libro['Ventas']['B5'].value.startswith('V-')
    assert libro['Detalle']['E5'].value == 'Kawasaki Ninja 400'


def test_el_reporte_no_esta_al_alcance_de_un_cliente(cliente_http, token_cliente):
    for ruta in ('', '/pdf', '/excel'):
        respuesta = cliente_http.get(
            f'/api/reportes/ventas-diarias{ruta}', headers=cabecera(token_cliente),
        )
        assert respuesta.status_code == 403, ruta


def test_el_empleado_si_puede_generar_el_reporte(cliente_http, token_empleado):
    respuesta = cliente_http.get('/api/reportes/ventas-diarias', headers=cabecera(token_empleado))

    assert respuesta.status_code == 200
